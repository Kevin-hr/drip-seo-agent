# -*- coding: utf-8 -*-
"""Build Pkgod-residual + P0-optimization Excel from scan data."""
import json, re, unicodedata, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"C:\Users\Administrator\Documents\drip-seo-agent\data\runs\v45-batch-2026-09-23\minimal-run"
all_scan = json.load(open(f"{BASE}\\all-products-scan.json", encoding="utf-8"))
clean = [x for x in all_scan if "error" not in x]

txt = open(r"D:\下单排行榜.txt", encoding="utf-8").read()
rank_raw = [l.strip() for l in txt.splitlines() if l.strip()]
rank_names = [re.sub(r"^\s*\d+\.\s*", "", l).strip() for l in rank_raw]

def norm(s):
    s = (s or "").lower()
    for a, b in [("&amp;", "&"), ("&#39;", "'"), ("&#8217;", "'"), ("’", "'"), ("&quot;", '"'), ("\u200b", "")]:
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()

rank_norm = [norm(n) for n in rank_names]

def match_rank(pname):
    pn = norm(pname)
    if not pn:
        return None
    if pn in rank_norm:
        i = rank_norm.index(pn)
        return (i + 1, "exact", rank_names[i])
    for i, rn in enumerate(rank_norm):
        if len(rn) >= 12 and pn.find(rn) >= 0:
            return (i + 1, "rank-in-product", rank_names[i])
    if len(pn) >= 12:
        for i, rn in enumerate(rank_norm):
            if rn.find(pn) >= 0:
                return (i + 1, "product-in-rank", rank_names[i])
    return None

def seo_status(p):
    has = [p.get("seoTitle"), p.get("seoKeyword"), p.get("seoDesc")]
    filled = sum(1 for v in has if v and v.strip())
    return filled, filled == 3

pkgod_rows = []
for p in clean:
    has_pkgod = "pkgod" in (p.get("urlValue") or "").lower() or "pkgod" in (p.get("url") or "").lower() or "pkgod" in (p.get("oldUrlValue") or "").lower()
    if not has_pkgod:
        continue
    m = match_rank(p.get("name") or "")
    filled, complete = seo_status(p)
    pkgod_rows.append({
        "id": p["id"], "name": (p.get("name") or "").strip(), "isShow": "是" if p.get("isShow") else "否",
        "urlValue": p.get("urlValue") or "", "url": p.get("url") or "", "oldUrlValue": p.get("oldUrlValue") or "",
        "seoTitle": (p.get("seoTitle") or "").strip(), "seoKeyword": (p.get("seoKeyword") or "").strip(), "seoDesc": (p.get("seoDesc") or "").strip(),
        "seoFilled": f"{filled}/3", "seoComplete": "是" if complete else "否",
        "p0": "P0" if m else "", "rank": m[0] if m else "", "rankName": m[2] if m else "", "matchType": m[1] if m else ""
    })
pkgod_rows.sort(key=lambda r: (r["p0"] != "P0", r["rank"] if r["rank"] else 99999))

# P0 list: ALL ranking-matched products with issues (pkgod residual OR incomplete SEO)
p0_rows = []
for p in clean:
    m = match_rank(p.get("name") or "")
    if not m:
        continue
    filled, complete = seo_status(p)
    has_pkgod = "pkgod" in (p.get("urlValue") or "").lower()
    issues = []
    if has_pkgod:
        issues.append("URL残留-Pkgod-")
    if not complete:
        missing = []
        for f in ["SeoTitle", "SeoKeyword", "SeoDesc"]:
            if not (p.get(f.lower()) or "").strip():
                missing.append(f)
        issues.append("SEO缺失:" + "+".join(missing))
    if not issues:
        continue
    p0_rows.append({
        "rank": m[0], "rankName": m[2], "id": p["id"], "name": (p.get("name") or "").strip(),
        "isShow": "是" if p.get("isShow") else "否", "urlValue": p.get("urlValue") or "", "url": p.get("url") or "",
        "issues": "; ".join(issues), "matchType": m[1]
    })
p0_rows.sort(key=lambda r: r["rank"])

wb = Workbook()
palette = {"header": "17365D", "second": "5B9BD5", "line": "B8C7D9", "p0": "C00000"}
line = Side(style="thin", color=palette["line"])

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = PatternFill("solid", fgColor=palette["header"])
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

def disp_w(v):
    return sum(2 if unicodedata.east_asian_width(ch) in ("F", "W") else 1 for ch in str(v or ""))

def fit_cols(ws, ncols, max_row, max_w=50):
    for c in range(1, ncols + 1):
        vals = [ws.cell(r, c).value for r in range(2, max_row + 1)]
        w = max((disp_w(v) for v in vals if v is not None), default=0) + 3
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(w + 3, max_w))

# Sheet 1: Pkgod residual
ws1 = wb.active
ws1.title = "Pkgod残留清单"
h1 = ["产品ID", "商品名称", "是否上架", "UrlValue", "前台URL", "旧UrlValue", "SeoTitle", "SeoKeyword", "SeoDesc", "SEO字段", "SEO完整", "P0标记", "排行榜排名", "排行榜商品名称", "匹配方式"]
ws1.append(h1)
for r in pkgod_rows:
    ws1.append([r["id"], r["name"], r["isShow"], r["urlValue"], r["url"], r["oldUrlValue"], r["seoTitle"], r["seoKeyword"], r["seoDesc"], r["seoFilled"], r["seoComplete"], r["p0"], r["rank"], r["rankName"], r["matchType"]])
style_header(ws1, len(h1))
for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
    for cell in row:
        cell.font = Font(name="微软雅黑", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
    if row[11].value == "P0":
        row[11].font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        row[11].fill = PatternFill("solid", fgColor=palette["p0"])
fit_cols(ws1, len(h1), ws1.max_row)

# Sheet 2: P0 optimization list
ws2 = wb.create_sheet("P0优化清单")
h2 = ["排行榜排名", "排行榜商品名称", "产品ID", "后台商品名称", "是否上架", "UrlValue", "前台URL", "需优化问题", "匹配方式"]
ws2.append(h2)
for r in p0_rows:
    ws2.append([r["rank"], r["rankName"], r["id"], r["name"], r["isShow"], r["urlValue"], r["url"], r["issues"], r["matchType"]])
style_header(ws2, len(h2))
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
    for cell in row:
        cell.font = Font(name="微软雅黑", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    row[7].font = Font(name="微软雅黑", size=10, bold=True, color=palette["p0"])
fit_cols(ws2, len(h2), ws2.max_row)

# Sheet 3: summary
ws3 = wb.create_sheet("汇总")
summary = [
    ["Drip Sneakers - URL/SEO 残留核查", ""],
    ["生成时间", "2026-09-24"],
    ["", ""],
    ["核查范围", "全部商品队列"],
    ["商品总数", len(clean)],
    ["残留 -Pkgod- URL 商品数", len(pkgod_rows)],
    ["其中命中下单排行榜（P0）", sum(1 for r in pkgod_rows if r["p0"])],
    ["P0 需优化商品总数（排行榜命中且 URL/SEO 有问题）", len(p0_rows)],
    ["", ""],
    ["说明", "P0 = 下单排行榜中的商品，其 URL 仍残留旧站 Pkgod 路径或 SEO 字段缺失，需优先修复为规范前台 URL 并补齐 SEO。"],
]
for row in summary:
    ws3.append(row)
ws3["A1"].font = Font(name="微软雅黑", size=14, bold=True)
ws3.column_dimensions["A"].width = 40
ws3.column_dimensions["B"].width = 60
for r in range(4, 10):
    ws3.cell(r, 1).font = Font(name="微软雅黑", size=10, bold=True)
    ws3.cell(r, 2).font = Font(name="微软雅黑", size=10)

out = r"C:\Users\Administrator\Documents\drip-seo-agent\output\Pkgod残留与P0优化清单.xlsx"
import os
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print(f"Saved {out}")
print(f"Pkgod rows: {len(pkgod_rows)} | P0 rows: {len(p0_rows)}")
json.dump({"pkgod": pkgod_rows, "p0": p0_rows, "summary": {"total": len(clean), "pkgod": len(pkgod_rows), "p0": len(p0_rows)}},
          open(f"{BASE}\\pkgod-excel-data.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)
